"""LipiScan backend API test suite.

Covers:
- /api/quality-check (clear vs blurry/dark/blank)
- /api/extract + /api/history (list + soft-delete)
- /api/append-docx (end + after_heading + heading-not-found)
- /api/download/{id} (valid docx + Bengali text survives)
"""
import io
import os
import zipfile
from pathlib import Path

import cv2
import numpy as np
import pytest
import requests
from docx import Document
from PIL import Image, ImageDraw, ImageFont

BASE_URL = os.environ["EXPO_PUBLIC_BACKEND_URL"].rstrip("/") if os.environ.get(
    "EXPO_PUBLIC_BACKEND_URL"
) else "https://docx-capture.preview.emergentagent.com"
API = f"{BASE_URL}/api"


# ---------- fixtures ----------
@pytest.fixture(scope="module")
def session():
    s = requests.Session()
    return s


def _find_font():
    for p in [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ]:
        if Path(p).exists():
            return p
    return None


def _make_text_image(text="Hello World OCR Test", size=(800, 400)) -> bytes:
    img = Image.new("RGB", size, "white")
    d = ImageDraw.Draw(img)
    font_path = _find_font()
    try:
        font = ImageFont.truetype(font_path, 56) if font_path else ImageFont.load_default()
    except Exception:
        font = ImageFont.load_default()
    d.text((40, 120), text, fill="black", font=font)
    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=95)
    return buf.getvalue()


def _make_blurry_image() -> bytes:
    data = _make_text_image("Hello World")
    arr = cv2.imdecode(np.frombuffer(data, np.uint8), cv2.IMREAD_COLOR)
    blurred = cv2.GaussianBlur(arr, (55, 55), 0)
    ok, buf = cv2.imencode(".jpg", blurred)
    return buf.tobytes()


def _make_dark_image() -> bytes:
    img = Image.new("RGB", (600, 400), (12, 12, 12))
    buf = io.BytesIO()
    img.save(buf, format="JPEG")
    return buf.getvalue()


def _make_blank_image() -> bytes:
    img = Image.new("RGB", (600, 400), "white")
    buf = io.BytesIO()
    img.save(buf, format="JPEG")
    return buf.getvalue()


def _make_docx(with_heading: bool = False) -> bytes:
    d = Document()
    d.add_paragraph("Intro paragraph.")
    if with_heading:
        d.add_heading("Section One", level=1)
        d.add_paragraph("Body under section one.")
        d.add_heading("Section Two", level=1)
        d.add_paragraph("Body under section two.")
    else:
        d.add_paragraph("Existing content line.")
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


# ---------- health ----------
def test_root(session):
    r = session.get(f"{API}/")
    assert r.status_code == 200
    assert r.json().get("message") == "LipiScan API"


# ---------- quality-check ----------
class TestQualityCheck:
    def test_clear_image_ok(self, session):
        r = session.post(
            f"{API}/quality-check",
            files={"file": ("clear.jpg", _make_text_image(), "image/jpeg")},
        )
        assert r.status_code == 200, r.text
        data = r.json()
        assert data["ok"] is True, f"issues: {data.get('issues')}"
        assert data["has_text"] is True
        assert data["issues"] == []
        assert data["blur_score"] > 45
        assert data["brightness"] > 55

    def test_blurry_image_has_blur_issue(self, session):
        r = session.post(
            f"{API}/quality-check",
            files={"file": ("blur.jpg", _make_blurry_image(), "image/jpeg")},
        )
        assert r.status_code == 200, r.text
        data = r.json()
        codes = {i["code"] for i in data["issues"]}
        assert "blur" in codes, f"expected blur, got {data}"
        assert data["ok"] is False

    def test_dark_image_has_dim_issue(self, session):
        r = session.post(
            f"{API}/quality-check",
            files={"file": ("dark.jpg", _make_dark_image(), "image/jpeg")},
        )
        assert r.status_code == 200
        data = r.json()
        codes = {i["code"] for i in data["issues"]}
        assert "dim" in codes
        assert data["ok"] is False

    def test_blank_image_has_no_text_issue(self, session):
        r = session.post(
            f"{API}/quality-check",
            files={"file": ("blank.jpg", _make_blank_image(), "image/jpeg")},
        )
        assert r.status_code == 200
        data = r.json()
        codes = {i["code"] for i in data["issues"]}
        assert "no_text" in codes
        assert data["ok"] is False

    def test_empty_file(self, session):
        r = session.post(
            f"{API}/quality-check",
            files={"file": ("empty.jpg", b"", "image/jpeg")},
        )
        assert r.status_code == 400


# ---------- extract + history ----------
class TestExtractAndHistory:
    created_ids = []

    def test_extract_creates_history(self, session):
        img = _make_text_image("TEST hello world")
        r = session.post(
            f"{API}/extract", files={"file": ("t.jpg", img, "image/jpeg")}
        )
        assert r.status_code == 200, r.text
        data = r.json()
        assert "id" in data and "text" in data
        assert isinstance(data["text"], str)
        TestExtractAndHistory.created_ids.append(data["id"])

    def test_history_lists_created(self, session):
        r = session.get(f"{API}/history")
        assert r.status_code == 200
        items = r.json()
        assert isinstance(items, list)
        ids = {i["id"] for i in items}
        for created in TestExtractAndHistory.created_ids:
            assert created in ids, f"created id {created} not in history"
        # sorted desc by created_at
        if len(items) >= 2:
            assert items[0]["created_at"] >= items[1]["created_at"]

    def test_history_soft_delete(self, session):
        assert TestExtractAndHistory.created_ids, "need at least one created id"
        target = TestExtractAndHistory.created_ids[0]
        r = session.delete(f"{API}/history/{target}")
        assert r.status_code == 200, r.text
        assert r.json().get("ok") is True
        # verify gone from GET
        r2 = session.get(f"{API}/history")
        ids = {i["id"] for i in r2.json()}
        assert target not in ids

    def test_history_delete_invalid_id(self, session):
        r = session.delete(f"{API}/history/not-an-oid")
        assert r.status_code == 400

    def test_history_delete_not_found(self, session):
        # valid-looking ObjectId that doesn't exist
        r = session.delete(f"{API}/history/507f1f77bcf86cd799439011")
        assert r.status_code == 404


# ---------- append-docx + download ----------
class TestAppendDocx:
    def test_append_end_and_download_contains_bengali(self, session):
        docx = _make_docx(with_heading=False)
        bn_text = "নতুন বাংলা লেখা"
        r = session.post(
            f"{API}/append-docx",
            data={"text": bn_text, "position": "end", "heading": ""},
            files={"file": ("in.docx", docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )
        assert r.status_code == 200, r.text
        payload = r.json()
        assert "download_id" in payload
        assert payload["filename"].endswith("-updated.docx")

        dl = session.get(f"{API}/download/{payload['download_id']}")
        assert dl.status_code == 200
        assert dl.headers.get("content-type", "").startswith(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
        # validate docx and Bengali
        d = Document(io.BytesIO(dl.content))
        paragraph_texts = [p.text for p in d.paragraphs]
        joined = "\n".join(paragraph_texts)
        assert "নতুন বাংলা" in joined, f"Bengali missing. Content: {joined!r}"
        # appended at end -> last non-empty paragraph should be the new one
        non_empty = [t for t in paragraph_texts if t.strip()]
        assert non_empty[-1].strip() == bn_text

    def test_append_after_heading_inserts_correctly(self, session):
        docx = _make_docx(with_heading=True)
        insert_text = "TEST inserted after section one"
        r = session.post(
            f"{API}/append-docx",
            data={"text": insert_text, "position": "after_heading", "heading": "Section One"},
            files={"file": ("h.docx", docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )
        assert r.status_code == 200, r.text
        dl = session.get(f"{API}/download/{r.json()['download_id']}")
        assert dl.status_code == 200
        d = Document(io.BytesIO(dl.content))
        texts = [p.text for p in d.paragraphs]
        # find heading index; new paragraph should be immediately after it
        idx = next(i for i, t in enumerate(texts) if "Section One" in t)
        assert texts[idx + 1].strip() == insert_text
        # ensure Section Two still exists after
        assert any("Section Two" in t for t in texts[idx + 2 :])

    def test_append_after_heading_not_found_returns_404(self, session):
        docx = _make_docx(with_heading=True)
        r = session.post(
            f"{API}/append-docx",
            data={"text": "abc", "position": "after_heading", "heading": "NonexistentHeading"},
            files={"file": ("h.docx", docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )
        assert r.status_code == 404
        assert "not found" in r.json().get("detail", "").lower()

    def test_append_invalid_docx_returns_400(self, session):
        r = session.post(
            f"{API}/append-docx",
            data={"text": "abc", "position": "end", "heading": ""},
            files={"file": ("bad.docx", b"not a docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )
        assert r.status_code == 400

    def test_download_missing_id_returns_404(self, session):
        r = session.get(f"{API}/download/does-not-exist")
        assert r.status_code == 404


# ---------- AI cleanup (new this iteration) ----------
class TestCleanup:
    def test_cleanup_mixed_bengali_english(self, session):
        raw = "Hell0 w0rld thls is a testt. আমি বাংলাय় লিখছ i"
        r = session.post(f"{API}/cleanup", json={"text": raw}, timeout=60)
        assert r.status_code == 200, r.text
        data = r.json()
        assert "text" in data
        cleaned = data["text"]
        assert isinstance(cleaned, str)
        assert cleaned.strip() != ""
        # sanity: should contain both Latin letters and Bengali script
        has_latin = any("a" <= c.lower() <= "z" for c in cleaned)
        has_bengali = any("\u0980" <= c <= "\u09FF" for c in cleaned)
        assert has_latin, f"expected English chars in cleaned output: {cleaned!r}"
        assert has_bengali, f"expected Bengali chars in cleaned output: {cleaned!r}"
        # should NOT contain obvious commentary preamble
        low = cleaned.lower()
        for banned in ["here is", "corrected text:", "sure,", "note:"]:
            assert banned not in low, f"model added commentary '{banned}': {cleaned!r}"

    def test_cleanup_empty_text_returns_400(self, session):
        r = session.post(f"{API}/cleanup", json={"text": ""})
        assert r.status_code == 400
        assert r.json().get("detail") == "No text to clean up"

    def test_cleanup_whitespace_only_returns_400(self, session):
        r = session.post(f"{API}/cleanup", json={"text": "   \n\t  "})
        assert r.status_code == 400
        assert r.json().get("detail") == "No text to clean up"

    def test_cleanup_missing_field_returns_422(self, session):
        r = session.post(f"{API}/cleanup", json={})
        # Pydantic validation -> 422
        assert r.status_code == 422
